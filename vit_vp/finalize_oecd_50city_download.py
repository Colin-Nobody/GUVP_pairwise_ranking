#!/usr/bin/env python3
"""Replace unreachable Mapillary IDs and integrity-check the final 10k image cohort."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import pandas as pd
from PIL import Image, ImageOps


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--oecd-xlsx", type=Path, required=True)
    args = parser.parse_args()
    root = args.root
    manifest_path = root / "images_manifest.csv"
    manifest = pd.read_csv(manifest_path, dtype={"uuid": str, "sequence_id": str, "oecd_region_code": str})
    download_log = pd.read_csv(root / "download_log.csv")
    failures = download_log.loc[download_log.status.eq("failed"), ["city", "orig_id", "detail"]].copy()
    rejected_path = root / "rejected_image_candidates.csv"
    corrupt_path = root / "corrupt_image_candidates.csv"
    rejected = pd.read_csv(rejected_path) if rejected_path.is_file() else (pd.read_csv(corrupt_path) if corrupt_path.is_file() else pd.DataFrame(columns=["uuid", "city", "orig_id", "error"]))
    if not rejected.empty:
        rejected_failures = rejected.rename(columns={"error": "detail"})[["city", "orig_id", "detail"]]
        rejected_failures = rejected_failures[~rejected_failures.orig_id.astype("int64").isin(failures.orig_id.astype("int64"))]
        failures = pd.concat([failures, rejected_failures], ignore_index=True)
    replacements = pd.read_csv(root / "replacement_manifest.csv", dtype={"uuid": str, "sequence_id": str, "oecd_region_code": str})
    initial_manifest_path = root / "images_manifest_initial_20261024.csv"
    initial_manifest = pd.read_csv(initial_manifest_path, dtype={"uuid": str, "sequence_id": str}) if initial_manifest_path.is_file() else manifest.copy()
    replacement_download = pd.concat([
        pd.read_csv(root / name) for name in [
            "copenhagen_replacement_download_log.csv",
            "cologne_replacement_download_log.csv",
            "vienna_replacement_download_log.csv",
            "additional_replacement_download_log.csv",
            "dedup_replacement_download_log.csv",
        ]
    ], ignore_index=True)
    replacement_ids = set(replacements.orig_id.astype("int64"))
    current_ids = set(manifest.orig_id.astype("int64"))
    replacement_rows = replacements[~replacements.orig_id.astype("int64").isin(current_ids)].copy()
    if len(failures) != len(replacement_rows):
        raise ValueError(f"Need exactly one replacement per final failed row: failures={len(failures)}, replacements={len(replacement_rows)}")
    if not replacement_download.status.isin(["downloaded", "skipped_existing"]).all():
        raise ValueError("At least one replacement download did not complete")

    substitution_rows = []
    final = manifest.copy()
    prior_replacements_path = root / "replacement_log.csv"
    prior = pd.read_csv(prior_replacements_path) if prior_replacements_path.is_file() else pd.DataFrame()
    if not prior.empty and {"failed_image_id", "replacement_uuid"}.issubset(prior.columns):
        for row in prior.itertuples(index=False):
            old = initial_manifest.loc[initial_manifest.orig_id.astype("int64").eq(int(row.failed_image_id))]
            if len(old) != 1:
                raise ValueError(f"Cannot recover original sample rank for replaced ID {row.failed_image_id}")
            replacement_hit = final.uuid.astype(str).eq(str(row.replacement_uuid))
            if replacement_hit.sum() != 1:
                raise ValueError(f"Previous replacement UUID is not unique in final manifest: {row.replacement_uuid}")
            final.loc[replacement_hit, "selection_rank_within_city"] = int(old.selection_rank_within_city.iloc[0])
            substitution_rows.append({
                "city": row.failed_city, "failed_mapillary_id": int(row.failed_image_id),
                "replacement_mapillary_id": int(row.replacement_image_id), "replacement_uuid": str(row.replacement_uuid),
                "sequence_id": str(row.replacement_sequence_id),
                "replaced_rank": int(old.selection_rank_within_city.iloc[0]),
                "download_failure": str(row.reason),
                "replacement_rule": "same city and OECD region, non-panoramic, unused sequence, 6 km urban-center radius, second fixed random seed",
            })
    for fail, replacement in zip(failures.sort_values(["city", "orig_id"]).itertuples(index=False),
                                 replacement_rows.sort_values(["city", "orig_id"]).itertuples(index=False)):
        if fail.city != replacement.city:
            raise ValueError(f"Replacement city mismatch: {fail.city} -> {replacement.city}")
        hit = final.orig_id.astype("int64").eq(int(fail.orig_id))
        if hit.sum() != 1:
            raise ValueError(f"Failed image ID is not unique in manifest: {fail.orig_id}")
        rank = int(final.loc[hit, "selection_rank_within_city"].iloc[0])
        final = final.loc[~hit].copy()
        replacement_row = replacement._asdict()
        replacement_row["selection_rank_within_city"] = rank
        final = pd.concat([final, pd.DataFrame([replacement_row])], ignore_index=True)
        substitution_rows.append({
            "city": fail.city, "failed_mapillary_id": int(fail.orig_id),
            "replacement_mapillary_id": int(replacement.orig_id), "replacement_uuid": replacement.uuid,
            "sequence_id": replacement.sequence_id, "replaced_rank": rank,
            "download_failure": str(fail.detail),
            "replacement_rule": "same city and OECD region, non-panoramic, unused sequence, 6 km urban-center radius, second fixed random seed",
        })
    final = final.sort_values(["city", "selection_rank_within_city"], kind="stable").reset_index(drop=True)
    if len(final) != 10_000 or final.uuid.nunique() != 10_000 or final.orig_id.nunique() != 10_000:
        raise ValueError("Final manifest must contain 10,000 unique rows and Mapillary IDs")
    if final.city.nunique() != 50 or not final.groupby("city").size().eq(200).all():
        raise ValueError("Final manifest must have 50 cities and 200 images per city")
    if not final.groupby("city").sequence_id.nunique().eq(200).all():
        raise ValueError("Final manifest has repeated image sequences within one or more cities")
    for city, group in final.groupby("city"):
        if sorted(group.selection_rank_within_city.astype(int)) != list(range(1, 201)):
            raise ValueError(f"Sampling ranks for {city} are not 1..200 after replacement")

    # Remove rejected corrupt files from the actual inference directory while
    # preserving their bytes and hashes in a separate audit quarantine.
    rejected_dir = root / "rejected_images"
    rejected_dir.mkdir(exist_ok=True)
    rejected_rows = []
    for row in rejected.itertuples(index=False):
        source = root / "images" / f"{row.uuid}.jpg"
        target = rejected_dir / source.name
        if source.is_file() and not target.exists():
            digest = sha256(source)
            source.replace(target)
        elif target.is_file():
            digest = sha256(target)
        else:
            digest = ""
        if target.is_file():
            rejected_rows.append({"uuid": row.uuid, "city": row.city, "mapillary_id": int(row.orig_id),
                                  "quarantine_path": str(target.resolve()), "bytes": target.stat().st_size,
                                  "sha256": digest, "reason": str(row.error)})

    ws = Image.open
    image_dir = root / "images"
    integrity = []
    for i, row in enumerate(final.itertuples(index=False), start=1):
        path = image_dir / f"{row.uuid}.jpg"
        if not path.is_file() or path.stat().st_size < 1000:
            raise FileNotFoundError(f"Missing or undersized final image: {path}")
        with ws(path) as im:
            im.verify()
        with ws(path) as im:
            oriented = ImageOps.exif_transpose(im)
            width, height = oriented.size
            fmt = im.format or ""
        if fmt != "JPEG":
            raise ValueError(f"Expected JPEG thumbnail for {row.uuid}, got {fmt}")
        integrity.append({"uuid": row.uuid, "city": row.city, "mapillary_id": int(row.orig_id),
                          "bytes": path.stat().st_size, "width": width, "height": height,
                          "format": fmt, "sha256": sha256(path), "decode": "ok"})
        if i % 1000 == 0:
            print(f"integrity {i}/10000 decoded and hashed", flush=True)
    integrity_df = pd.DataFrame(integrity)
    if integrity_df.sha256.nunique() != 10_000:
        duplicates = integrity_df[integrity_df.sha256.duplicated(False)]
        duplicates.to_csv(root / "duplicate_content_candidates.csv", index=False)
        raise ValueError(f"Found {10_000 - integrity_df.sha256.nunique()} duplicate image contents")
    actual_jpgs = list(image_dir.glob("*.jpg"))
    if len(actual_jpgs) != 10_000:
        raise ValueError(f"Expected exactly 10,000 final JPEGs in images/, found {len(actual_jpgs)}")

    wb_manifest = final[["city", "country", "oecd_region", "oecd_region_code"]].drop_duplicates()
    from openpyxl import load_workbook
    wb = load_workbook(args.oecd_xlsx, read_only=True, data_only=True)
    rows = list(wb["欧洲地区_11主题分数"].iter_rows(values_only=True))
    header_index = next(i for i, r in enumerate(rows) if r and r[0] == "国家")
    refs = {(str(r[0]).strip(), str(r[1]).strip(), str(r[2]).strip())
            for r in rows[header_index + 1:] if r and r[0] and r[1] and r[2]}
    # The workbook has country/region/code keys; city is only our local label.
    bad = [tuple(map(str, x)) for x in wb_manifest[["country", "oecd_region", "oecd_region_code"]].itertuples(index=False, name=None)
           if tuple(map(str, x)) not in refs]
    if wb_manifest.city.nunique() != 50 or wb_manifest.oecd_region_code.nunique() != 50 or bad:
        raise ValueError(f"OECD city/region matches are not 50 unique exact rows: {bad}")

    all_substitutions = substitution_rows
    if all_substitutions:
        all_substitutions = list({int(row["failed_mapillary_id"]): row for row in all_substitutions}.values())

    final.to_csv(manifest_path, index=False, encoding="utf-8-sig")
    pd.DataFrame(all_substitutions).to_csv(root / "replacement_log.csv", index=False, encoding="utf-8-sig")
    integrity_df.to_csv(root / "image_integrity.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(rejected_rows).to_csv(root / "rejected_image_log.csv", index=False, encoding="utf-8-sig")
    audit = {
        "final_images": len(final), "cities": int(final.city.nunique()), "images_per_city": 200,
        "unique_mapillary_ids": int(final.orig_id.nunique()), "unique_sequences": int(final.sequence_id.nunique()),
        "unique_content_sha256": int(integrity_df.sha256.nunique()), "decoded_images": int(len(integrity_df)),
        "oecd_exact_region_matches": int(wb_manifest.oecd_region_code.nunique()),
        "replacement_count": len(all_substitutions), "rejected_images_quarantined": len(rejected_rows),
        "substitutions": all_substitutions,
        "image_resolution": "Mapillary thumb_1024 when accessible; thumb_original fallback when transformed URL returns 4xx; all inputs use the exact V8 512-long-edge resize at scoring time",
        "storage": str(root.resolve()),
    }
    (root / "download_audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: audit[k] for k in ["final_images", "cities", "unique_mapillary_ids", "unique_sequences", "unique_content_sha256", "decoded_images", "oecd_exact_region_matches", "replacement_count"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
